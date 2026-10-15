"""
Unit tests for the monthly consumption report tool.

Ported from vivaldi_webai's src/lib/consumos/report.ts: the trustworthy
number is the invoiced sales (account.move.line, posted out_invoice) for a
single Odoo company, per client, for a given YYYY-MM period - never a
sale.order total, which mixes companies and non-invoiced drafts.

The client list itself is NOT hardcoded: it's resolved live from Odoo's
native res.partner.category tags (e.g. "Avendra - Reporte Consumo
Vivaldi"), so Oscar/Yossymar can add or remove a hotel from the report
group directly in Odoo's contact form - no redeploy needed.

Run with: pytest tests/unit/test_consumption_report.py -v -m unit
"""

import base64
from unittest.mock import AsyncMock

import openpyxl
import pytest

pytestmark = [pytest.mark.unit]


class TestPeriodBounds:
    def test_first_and_last_day_of_month(self):
        from odoo_mcp_server.tools.consumption_report import period_bounds

        assert period_bounds("2026-08") == ("2026-08-01", "2026-08-31")

    def test_leap_february(self):
        from odoo_mcp_server.tools.consumption_report import period_bounds

        assert period_bounds("2028-02") == ("2028-02-01", "2028-02-29")

    def test_rejects_malformed_period(self):
        from odoo_mcp_server.odoo.exceptions import OdooValidationError
        from odoo_mcp_server.tools.consumption_report import period_bounds

        with pytest.raises(OdooValidationError):
            period_bounds("agosto-2026")


class TestBuildConsumptionDomain:
    def test_domain_scopes_to_invoiced_product_lines_one_company(self):
        from odoo_mcp_server.tools.consumption_report import build_consumption_domain

        domain = build_consumption_domain(
            commercial_partner_ids=[8727, 4521],
            inicio="2026-08-01",
            fin="2026-08-31",
            company_id=1,
        )

        assert ["move_id.move_type", "=", "out_invoice"] in domain
        assert ["move_id.state", "=", "posted"] in domain
        assert ["move_id.invoice_date", ">=", "2026-08-01"] in domain
        assert ["move_id.invoice_date", "<=", "2026-08-31"] in domain
        assert ["partner_id.commercial_partner_id", "in", [8727, 4521]] in domain
        assert ["display_type", "=", "product"] in domain
        assert ["company_id", "=", 1] in domain


class TestResolvePartnersByTag:
    @pytest.fixture
    def mock_odoo_client(self):
        return AsyncMock()

    @pytest.mark.asyncio
    async def test_looks_up_tag_then_its_partners(self, mock_odoo_client):
        from odoo_mcp_server.tools.consumption_report import resolve_partners_by_tag

        mock_odoo_client.search_read.side_effect = [
            [{"id": 12}],  # res.partner.category lookup
            [{"id": 8727, "name": "HAMACA BEACH RESORT SAS", "vat": "101172381"}],
        ]

        partners = await resolve_partners_by_tag(mock_odoo_client, "Avendra - Reporte Consumo Vivaldi")

        assert partners == [{"id": 8727, "name": "HAMACA BEACH RESORT SAS", "vat": "101172381"}]
        first_call = mock_odoo_client.search_read.await_args_list[0]
        assert first_call.kwargs["model"] == "res.partner.category"
        second_call = mock_odoo_client.search_read.await_args_list[1]
        assert second_call.kwargs["domain"] == [["category_id", "in", [12]]]

    @pytest.mark.asyncio
    async def test_raises_clear_error_when_tag_missing(self, mock_odoo_client):
        from odoo_mcp_server.odoo.exceptions import OdooValidationError
        from odoo_mcp_server.tools.consumption_report import resolve_partners_by_tag

        mock_odoo_client.search_read.return_value = []

        with pytest.raises(OdooValidationError):
            await resolve_partners_by_tag(mock_odoo_client, "Tag que no existe")


class TestGenerateConsumptionReportTool:
    @pytest.fixture
    def mock_odoo_client(self):
        return AsyncMock()

    @pytest.fixture
    def report_groups(self):
        return {
            "avendra": "Avendra - Reporte Consumo Vivaldi",
            "quantum": "Quantum - Reporte Consumo Vivaldi",
        }

    def _client_with_tag_and_invoices(self, mock_odoo_client, partners, groups, commercial_of=None):
        """groups: read_group rows as Odoo returns them, e.g.
        {"partner_id": [8727, "X"], "__count": 2, "price_subtotal": 36108.0}.
        commercial_of: partner id -> commercial partner id (default: itself)."""
        commercial_of = commercial_of or {}

        async def fake_search_read(model, domain=None, fields=None, limit=None, order=None):
            if model == "res.partner.category":
                return [{"id": 12}]
            if model == "res.partner":
                return partners
            raise AssertionError(f"unexpected search_read on {model}")

        async def fake_read(model, ids, fields=None):
            assert model == "res.partner"
            return [
                {"id": i, "commercial_partner_id": [commercial_of.get(i, i), "x"]} for i in ids
            ]

        mock_odoo_client.search_read.side_effect = fake_search_read
        mock_odoo_client.read.side_effect = fake_read
        mock_odoo_client.read_group.return_value = groups

    @pytest.mark.asyncio
    async def test_reports_consumo_per_client_and_builds_xlsx(self, mock_odoo_client, report_groups):
        from odoo_mcp_server.tools.consumption_report import execute_tool

        partners = [
            {"id": 8727, "name": "HAMACA BEACH RESORT SAS", "vat": "101172381"},
            {"id": 4521, "name": "VIVA MICHES SAS", "vat": "131624138"},
        ]
        self._client_with_tag_and_invoices(
            mock_odoo_client,
            partners,
            [{"partner_id": [8727, "HAMACA"], "__count": 2, "price_subtotal": 36108.0}],
        )

        result = await execute_tool(
            "generate_consumption_report",
            {"periodo": "2026-08", "grupo": "avendra"},
            mock_odoo_client,
            report_groups=report_groups,
            company_id=1,
        )

        import json

        payload = json.loads(result[0].text)
        assert payload["periodo"] == "2026-08"
        assert payload["grupo"] == "avendra"
        hamaca = next(r for r in payload["rows"] if r["rnc"] == "101172381")
        assert hamaca["consumoDop"] == pytest.approx(36108.0)
        assert hamaca["facturas"] == 2
        viva_miches = next(r for r in payload["rows"] if r["rnc"] == "131624138")
        assert viva_miches["consumoDop"] == 0
        assert payload["totalDop"] == pytest.approx(36108.0)

        # The xlsx must actually be a valid, readable workbook with the right numbers.
        xlsx_bytes = base64.b64decode(payload["xlsx_base64"])
        wb = openpyxl.load_workbook(filename=__import__("io").BytesIO(xlsx_bytes))
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert "RNC" in header
        values = [tuple(row) for row in ws.iter_rows(min_row=2, values_only=True)]
        assert any(row[header.index("RNC")] == "101172381" for row in values)

    @pytest.mark.asyncio
    async def test_missing_grupo_lists_valid_options(self, mock_odoo_client, report_groups):
        from odoo_mcp_server.odoo.exceptions import OdooValidationError
        from odoo_mcp_server.tools.consumption_report import execute_tool

        with pytest.raises(OdooValidationError) as exc_info:
            await execute_tool(
                "generate_consumption_report",
                {"periodo": "2026-08"},
                mock_odoo_client,
                report_groups=report_groups,
                company_id=1,
            )

        assert "avendra" in str(exc_info.value)
        assert "quantum" in str(exc_info.value)

    @pytest.mark.asyncio
    async def test_single_configured_group_used_as_default(self, mock_odoo_client):
        from odoo_mcp_server.tools.consumption_report import execute_tool

        self._client_with_tag_and_invoices(mock_odoo_client, [], [])

        result = await execute_tool(
            "generate_consumption_report",
            {"periodo": "2026-08"},
            mock_odoo_client,
            report_groups={"avendra": "Avendra - Reporte Consumo Vivaldi"},
            company_id=1,
        )

        import json

        assert json.loads(result[0].text)["grupo"] == "avendra"

    @pytest.mark.asyncio
    async def test_unknown_grupo_rejected(self, mock_odoo_client, report_groups):
        from odoo_mcp_server.odoo.exceptions import OdooValidationError
        from odoo_mcp_server.tools.consumption_report import execute_tool

        with pytest.raises(OdooValidationError):
            await execute_tool(
                "generate_consumption_report",
                {"periodo": "2026-08", "grupo": "inventado"},
                mock_odoo_client,
                report_groups=report_groups,
                company_id=1,
            )

    @pytest.mark.asyncio
    async def test_no_groups_configured_raises_clear_error(self, mock_odoo_client):
        from odoo_mcp_server.odoo.exceptions import OdooValidationError
        from odoo_mcp_server.tools.consumption_report import execute_tool

        with pytest.raises(OdooValidationError):
            await execute_tool(
                "generate_consumption_report",
                {"periodo": "2026-08"},
                mock_odoo_client,
                report_groups={},
                company_id=1,
            )

    @pytest.mark.asyncio
    async def test_single_read_group_call_for_all_clients(self, mock_odoo_client, report_groups):
        """N clients must cost ONE aggregate query, not N line-level searches."""
        from odoo_mcp_server.tools.consumption_report import execute_tool

        partners = [{"id": i, "name": f"H{i}", "vat": str(i)} for i in (1, 2, 3, 4, 5, 6)]
        self._client_with_tag_and_invoices(mock_odoo_client, partners, [])

        await execute_tool(
            "generate_consumption_report",
            {"periodo": "2026-08", "grupo": "avendra"},
            mock_odoo_client,
            report_groups=report_groups,
            company_id=1,
        )

        assert mock_odoo_client.read_group.await_count == 1
        domain = mock_odoo_client.read_group.await_args.kwargs["domain"]
        assert ["partner_id.commercial_partner_id", "in", [1, 2, 3, 4, 5, 6]] in domain

    @pytest.mark.asyncio
    async def test_child_contact_lines_roll_up_to_commercial_partner(self, mock_odoo_client, report_groups):
        """An invoice addressed to a child contact (e.g. delivery address)
        must count toward its commercial partner's total."""
        import json

        from odoo_mcp_server.tools.consumption_report import execute_tool

        partners = [{"id": 8727, "name": "HAMACA", "vat": "101172381"}]
        self._client_with_tag_and_invoices(
            mock_odoo_client,
            partners,
            [
                {"partner_id": [8727, "HAMACA"], "__count": 1, "price_subtotal": 100.0},
                {"partner_id": [9999, "HAMACA, Almacen"], "__count": 2, "price_subtotal": 50.0},
            ],
            commercial_of={9999: 8727},
        )

        result = await execute_tool(
            "generate_consumption_report",
            {"periodo": "2026-08", "grupo": "avendra"},
            mock_odoo_client,
            report_groups=report_groups,
            company_id=1,
        )

        row = json.loads(result[0].text)["rows"][0]
        assert row["consumoDop"] == pytest.approx(150.0)
        assert row["facturas"] == 3

    @pytest.mark.asyncio
    async def test_no_partners_in_tag_skips_aggregate_query(self, mock_odoo_client, report_groups):
        from odoo_mcp_server.tools.consumption_report import execute_tool

        self._client_with_tag_and_invoices(mock_odoo_client, [], [])

        await execute_tool(
            "generate_consumption_report",
            {"periodo": "2026-08", "grupo": "avendra"},
            mock_odoo_client,
            report_groups=report_groups,
            company_id=1,
        )

        mock_odoo_client.read_group.assert_not_called()
